import re, difflib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC = "/root/.claude/uploads/a7ca4361-b1fd-50b0-bc6b-9819d20c31ff/a975f3b8-Rubriques_Hierarchie_De_Gestion_1.pdf"
OUT = "/home/user/pigment/Mapping_Rubriques_Hierarchie_Gestion.xlsx"

# ---------- 1. Parse the PDF: level = left indentation of each line ----------
import pymupdf
pdf = pymupdf.open(SRC)
items = []  # [code, label, depth]
for p in pdf:
    for b in p.get_text("dict")["blocks"]:
        for l in b.get("lines", []):
            t = "".join(s["text"] for s in l["spans"]).strip()
            x = l["bbox"][0]
            if not t or x < 10 or x > 500 or re.match(r"^lundi \d+ mai 2026", t):
                continue  # page footer (date / page number)
            m = re.match(r"^([A-Z0-9][A-Za-z0-9_.\-]*) - (.*)$", t)
            if m:
                items.append([m.group(1), m.group(2), round((x - 17.9) / 28.34)])
            else:  # label wrapped on the next line
                items[-1][1] += " " + t

# ---------- 2. Build the tree ----------
stack, rows = [], []
for i, (code, label, dep) in enumerate(items):
    while stack and stack[-1][0] >= dep:
        stack.pop()
    assert (stack[-1][0] if stack else -1) == dep - 1, (code, dep)
    parent = stack[-1][1] if stack else ""
    is_node = i + 1 < len(items) and items[i + 1][2] > dep
    rows.append(dict(code=code, label=label, depth=dep, parent=parent, node=is_node,
                     path=[c for _, c in stack] + [code]))
    stack.append((dep, code))
assert len({r["code"] for r in rows}) == len(rows)

# ---------- 3. Labels ----------
SPLIT_FIX = {  # manual short-label cut where the automatic split is wrong
    "A35109": "dep GW stés MEE", "A6436": "Dériv act.JV - BFR(chge) - CFH", "A6437": "Dériv act.JV - BFR(chge) - FVH",
    "A6438": "Dériv act. JV - BFR(chge) - NH", "P6436": "Dériv pas.JV - BFR(chge) - CFH", "P6437": "Dériv pas.JV- BFR(chge) - FVH",
    "P6438": "Dériv pas. JV - BFR(chge) - NH", "A70109": "Écrêtement - Net IDA - C", "P10": "Capitaux propres - Gpe",
    "R7220": "ID régul N-1", "A15109": "dep GW", "XIA315BN": "Diff. temporelles \"régul\"",
}
def split_label(t, code=None):
    if code in SPLIT_FIX and t.startswith(SPLIT_FIX[code] + " - "):
        return SPLIT_FIX[code], t[len(SPLIT_FIX[code]) + 3:]
    parts = t.split(" - ")
    if len(parts) == 1:
        return t, t
    best = None
    for i in range(1, len(parts)):
        s, l = " - ".join(parts[:i]), " - ".join(parts[i:])
        score = difflib.SequenceMatcher(None, s.lower(), l.lower()[:max(len(s), 1) + 5]).ratio()
        if len(s) > 30: score -= 1  # short labels are max 30 chars in the source
        if s.count("(") != s.count(")") or l.count("(") != l.count(")"): score -= 1
        if best is None or score > best[0]:
            best = (score, s, l)
    return best[1], best[2]

# ---------- 4. Excel ----------
HDR = Font(bold=True, color="FFFFFF"); FILL = PatternFill("solid", fgColor="1F4E78")
def header(ws, cols):
    ws.append(cols)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HDR, FILL, Alignment(vertical="center")
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x

maxd = max(r["depth"] for r in rows)
wb = Workbook()

# Sheet 1: flattened mapping, one row per leaf
ws = wb.active; ws.title = "Mapping_Niveaux"
cols = [f"Niveau {i}" for i in range(maxd + 1)] + ["Code feuille", "Profondeur"]
header(ws, cols)
for r in rows:
    if r["node"]: continue
    ws.append(r["path"] + [""] * (maxd + 1 - len(r["path"])) + [r["code"], r["depth"]])
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
widths(ws, [12] * (maxd + 1) + [14, 11])

# Sheet 2: parent / child
ws = wb.create_sheet("Parent_Enfant")
header(ws, ["Ordre", "Code", "Code parent", "Niveau", "Type"])
for i, r in enumerate(rows, 1):
    ws.append([i, r["code"], r["parent"], r["depth"], "Noeud" if r["node"] else "Feuille"])
ws.auto_filter.ref = f"A1:E{ws.max_row}"
widths(ws, [8, 14, 14, 8, 10])

# Sheet 3: labels (for Pigment upload)
ws = wb.create_sheet("Libelles")
header(ws, ["Code", "Libellé court", "Libellé long", "Code - Libellé", "Libellé source complet"])
for r in rows:
    s, l = split_label(r["label"], r["code"])
    ws.append([r["code"], s, l, f'{r["code"]} - {l}', r["label"]])
ws.auto_filter.ref = f"A1:E{ws.max_row}"
widths(ws, [14, 32, 70, 80, 90])

# Sheet 4: one list per level (Pigment dimension lists)
ws = wb.create_sheet("Listes_par_niveau")
header(ws, ["Niveau", "Code", "Code parent", "Libellé long"])
for d in range(maxd + 1):
    for r in rows:
        if r["depth"] == d:
            ws.append([d, r["code"], r["parent"], split_label(r["label"], r["code"])[1]])
ws.auto_filter.ref = f"A1:D{ws.max_row}"
widths(ws, [8, 14, 14, 80])

# Sheet 5: readme
ws = wb.create_sheet("Lisez-moi")
notes = [
    f"Source : Rubriques_Hierarchie_De_Gestion (PDF, {len(pdf)} pages, {len(rows)} rubriques).",
    "Niveaux lus directement depuis l'indentation du PDF (Niveau 0 = rubrique racine : B1, B2, ..., P1, R, XTF, ZIMPOT, ...).",
    "Logique : chaque rubrique est rattachée à la rubrique moins indentée qui la précède (ex. A10108 > A1008 > A10 > A10R > B1).",
    "Mapping_Niveaux : 1 ligne par rubrique feuille avec toute la chaîne de ses parents. Hiérarchie irrégulière : colonnes vides au-delà de la profondeur de la feuille.",
    "Parent_Enfant : format parent/enfant (recommandé pour Pigment). Libelles : code + libellés pour l'import. Listes_par_niveau : une liste par niveau avec son parent.",
    "Libellé court / long : découpage automatique du texte source 'court - long' (le libellé source complet est conservé en colonne E).",
    f"Profondeur max : {maxd}. Nœuds : {sum(r['node'] for r in rows)}. Feuilles : {sum(not r['node'] for r in rows)}.",
]
for n in notes: ws.append([n])
ws.column_dimensions["A"].width = 160
wb.move_sheet("Lisez-moi", offset=-4)
wb.save(OUT)
print("OK", OUT, len(rows), "maxdepth", maxd, notes[-1])
